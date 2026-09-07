import io
import fitz
from docx import Document
from fastapi import APIRouter, Response, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse

router = APIRouter(prefix="/download", tags=["Document Downloads"])


@router.get("/sim-policy")
async def get_sim_policy_html():
    """Downloadable Policy 1: SIM & eSIM Replacement (HTML)."""
    html_content = """
    <!DOCTYPE html>
    <html>
    <head><title>SIM Card & eSIM Replacement Policy</title></head>
    <body style="font-family: sans-serif; padding: 30px; line-height: 1.6;">
        <h1>Telecom Policy: Lost & Damaged SIM Card Replacement</h1>
        <p><strong>Ref:</strong> POL-SIM-2026 | <strong>Effective:</strong> Jan 2026</p>
        <hr/>
        <h2>1. Instant Free eSIM Activation</h2>
        <p>Customers with eSIM devices can activate a digital eSIM QR-code instantly for $0.00 within 5 minutes over Wi-Fi.</p>
        <h2>2. Physical SIM Replacement</h2>
        <p>Physical SIM cards are available at all stores. The standard $5 fee is waived for VIP and Gold accounts.</p>
        <h2>3. Stolen Device Line Bar</h2>
        <p>Immediate outgoing call/data bar is placed upon reporting stolen phones to protect customers from bill shock.</p>
    </body>
    </html>
    """
    return HTMLResponse(content=html_content)


@router.get("/wifi-policy")
async def download_wifi_policy_pdf():
    """Downloadable Policy 2: WiFi Troubleshooting & Dispatch Policy (PDF)."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)

    page.insert_text((50, 60), "TELECOM SUPPORT POLICY: HOME WIFI & BROADBAND OUTAGES", fontsize=13, fontname="helv", color=(0.1, 0.2, 0.5))
    page.insert_text((50, 80), "Document ID: POL-WIFI-TECH-2026 | Effective Date: January 2026", fontsize=9, fontname="helv", color=(0.3, 0.3, 0.3))

    t1 = (
        "1. Red Light & No Internet Diagnostic Procedure\n"
        "When a residential customer reports a red light on their router:\n"
        "- Step 1: Perform 30-second power cycle (unplug power cable, wait 30 seconds, plug back in).\n"
        "- Step 2: Ensure the green optical cable clicked firmly into the ONT box.\n"
        "- Step 3: CSR performs a remote line diagnostic check from CRM console."
    )
    page.insert_textbox(fitz.Rect(50, 110, 545, 210), t1, fontsize=9.5, fontname="helv")

    t2 = (
        "2. Free On-Site Technician Dispatch Guarantee\n"
        "If power cycling does not restore internet connectivity:\n"
        "- Schedule an Urgent Free On-Site Technician Dispatch.\n"
        "- Dispatch SLA: Within 4 hours for Priority/Work-From-Home accounts, within 24 hours for standard accounts.\n"
        "- Visit Fee: $0.00 (100% Free for network/hardware faults)."
    )
    page.insert_textbox(fitz.Rect(50, 230, 545, 330), t2, fontsize=9.5, fontname="helv")

    t3 = (
        "3. Emergency 5G Backup Mobile Data\n"
        "During WiFi downtime, CSR agents are authorized to add 50GB of free emergency 5G mobile data\n"
        "to the customer's phone while waiting for the technician to complete the repair."
    )
    page.insert_textbox(fitz.Rect(50, 350, 545, 430), t3, fontsize=9.5, fontname="helv")

    pdf_bytes = doc.tobytes()
    doc.close()

    return StreamingResponse(
        io.BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=wifi_troubleshooting_policy.pdf"}
    )


@router.get("/billing-policy")
async def download_billing_policy_docx():
    """Downloadable Policy 3: Billing Dispute & Refund Policy (DOCX)."""
    doc = Document()
    doc.add_heading("Customer Service Policy: Billing Disputes & Late Fee Waivers", level=1)
    doc.add_paragraph("Document ID: POL-BILL-REFUND-2026 | Department: Billing & Retention")

    doc.add_heading("1. One-Time Annual Courtesy Late Fee Waiver", level=2)
    doc.add_paragraph(
        "Any active customer charged a $15.00 late payment fee is eligible for an automatic one-time waiver "
        "if no other late fee waiver was granted in the last 12 months. The $15.00 credit is applied to the next invoice."
    )

    doc.add_heading("2. Erroneous Third-Party Charges", level=2)
    doc.add_paragraph(
        "CSRs can instantly issue refunds up to $50.00 for accidental third-party micro-transactions."
    )

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=billing_refund_policy.docx"}
    )
