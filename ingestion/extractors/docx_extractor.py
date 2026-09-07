import os
from docx import Document
from core.logger import get_logger
from ingestion.base import BaseExtractor, ExtractedDocument

logger = get_logger()


class DocxExtractor(BaseExtractor):
    """
    Extracts structured policy rules, heading hierarchies, tables, and clauses
    from Word (.docx) documents.
    Ideal for Billing Dispute and Fee Refund Policies.
    """

    def extract(self, file_path: str) -> ExtractedDocument:
        logger.info(f"Extracting DOCX document from: {file_path}")
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        doc = Document(file_path)
        sections = []
        current_title = "Billing Dispute & Refund Policy"
        current_paragraphs = []

        # Process paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if p.style.name.startswith("Heading"):
                if current_paragraphs:
                    sections.append({
                        "section_title": current_title,
                        "content": "\n".join(current_paragraphs),
                        "page_number": 1
                    })
                    current_paragraphs = []
                current_title = text
            else:
                current_paragraphs.append(text)

        if current_paragraphs:
            sections.append({
                "section_title": current_title,
                "content": "\n".join(current_paragraphs),
                "page_number": 1
            })

        # Process tables
        for idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            table_text = "\n".join(table_rows)
            sections.append({
                "section_title": f"Policy Refund Schedule Table {idx + 1}",
                "content": table_text,
                "page_number": 1
            })

        full_text = "\n\n".join([f"### {s['section_title']}\n{s['content']}" for s in sections])

        return ExtractedDocument(
            document_id=f"doc_docx_{os.path.basename(file_path)}",
            source_path=file_path,
            title="Billing Fee Refund & Courtesy Credit Policy",
            raw_text=full_text,
            policy_type="billing_refund",
            sections=sections,
            metadata={"format": "docx", "total_sections": len(sections)}
        )
